#!/usr/bin/env python3
"""Genera assets/estudios-data.js.

Dos fuentes:

* ``estudios/Top 100 estudios.xlsx`` — nombre, precio, entrega, preparación
  y sucursales. Es la autoridad sobre los datos que se muestran.
* ``estudios/diccionario-busqueda.md`` — el diccionario maestro de búsqueda:
  alias directos, grupos semánticos y errores de escritura frecuentes por
  estudio, más la tabla global de erratas. Solo alimenta al buscador; nunca
  cambia lo que ve el paciente.

Las curadurías del sitio (categoría corregida, destacados, campaña de la
mujer, precios "a consultar") viven en las tablas de abajo.

Uso:  python3 tools/build_estudios.py
"""
import io
import json
import os
import re
import unicodedata

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, 'estudios', 'Top 100 estudios.xlsx')
DICT = os.path.join(ROOT, 'estudios', 'diccionario-busqueda.md')
OUT = os.path.join(ROOT, 'assets', 'estudios-data.js')

# Columna1 del Excel -> categoría interna / etiqueta pública / icono lucide
CATEGORIES = {
    'Laboratorio': ('Laboratorio (procesar)', 'Laboratorio clínico', 'test-tube-diagonal'),
    'Rayos X': ('Rayos X', 'Rayos X', 'scan-line'),
    'USG': ('Ultrasonidos', 'Ultrasonido', 'radar'),
    'Tomografia': ('Tomografía', 'Tomografía', 'scan'),
    'Especiales': ('Estudios especiales', 'Estudios especiales', 'sparkles'),
}

# Correcciones de categoría sobre el Excel (errores de captura en la fuente).
CATEGORY_FIX = {
    'GLU.S': 'Laboratorio',   # glucosa en suero es laboratorio, no rayos X
}

# Sucursal (Excel) -> unidades reales
BRANCHES = {
    'A B & C': ['Tultepec', 'Cantú', 'Tepalcapa', 'Haciendas', 'Joya', 'Tepojaco'],
    'A & B': ['Tultepec', 'Cantú', 'Tepalcapa', 'Haciendas', 'Joya'],
    'A': ['Tultepec', 'Cantú'],
}

TOP = {'QSC-6', 'EGO', 'BHC-FEM', 'BHC-MASC', 'PB5-38', 'PB4-27', 'PB2-12E', 'TX-1', 'COLPO'}
ALTA = {'MASTO', 'TAC.CRANS', 'TAC-ABDIC', 'TAC-ABDSC', 'PAPS', 'COLPO-VAG'}
CAMPAIGN = {'COLPO', 'PAPS', 'COLPO-VAG'}

# Sinónimos extra por estudio, además de los que trae la columna "Descripción".
EXTRA_KEYWORDS = {
    'EGO': ['orina', 'examen de orina', 'pipi'],
    'HD-G': ['diabetes', 'azucar', 'hba1c', 'control de diabetes'],
    'GLU.S': ['azucar en sangre', 'diabetes', 'glucemia'],
    'PB5-38': ['colesterol', 'trigliceridos', 'acido urico', 'perfil de lipidos'],
    'PB4-27': ['colesterol', 'trigliceridos', 'acido urico', 'perfil de lipidos'],
    'PB2-12E': ['colesterol', 'trigliceridos', 'acido urico'],
    'QSC-6': ['colesterol', 'trigliceridos', 'acido urico', 'urea'],
    'BHC-FEM': ['anemia', 'plaquetas', 'globulos rojos', 'globulos blancos'],
    'BHC-MASC': ['anemia', 'plaquetas', 'globulos rojos', 'globulos blancos'],
    'BHC1A6A': ['anemia en niños', 'pediatrico', 'bebe'],
    'BHC6A12A': ['anemia en niños', 'pediatrico', 'escolar'],
    'MASTO': ['mama', 'seno', 'senos', 'pecho', 'cancer de mama', 'mamografia'],
    'USG-GM': ['mama', 'seno', 'senos', 'pecho', 'bolita en el seno'],
    'HIV': ['sida', 'vih', 'prueba de vih'],
    'DUO COV-INF': ['covid', 'coronavirus', 'gripe', 'influenza', 'prueba covid'],
    'CPL': ['parasitos', 'heces', 'excremento', 'popo', 'diarrea'],
    'SOH': ['heces', 'sangre en el excremento', 'colon'],
    'SAL-VISU': ['ojos', 'vista', 'lentes', 'graduacion'],
    'DENCOLCA': ['huesos', 'osteoporosis', 'densidad osea'],
    'PDA-1': ['antidoping', 'drogas', 'examen de drogas'],
    'PDA-2': ['antidoping', 'drogas', 'examen de drogas'],
    'PRENUP': ['boda', 'matrimonio', 'certificado prenupcial'],
    'IR-HOMA': ['resistencia a la insulina', 'prediabetes'],
    'INSULI': ['resistencia a la insulina', 'prediabetes'],
    'VITA-D': ['vitamina d', 'deficiencia de vitaminas'],
    'ECO': ['corazon', 'ultrasonido del corazon'],
    'EKG': ['corazon', 'electro'],
    'EKG/CARDIOLO': ['corazon', 'electro'],
    'EKG-ESF': ['corazon', 'prueba de esfuerzo'],
}

# Precios no numéricos del Excel -> etiqueta que se muestra en la tarjeta
PRICE_LABELS = {
    'Depende del modelo': 'Depende del modelo',
}


def clean(text):
    """Colapsa espacios dobles y recorta."""
    return re.sub(r'\s+', ' ', str(text or '')).strip()


def strip_accents(text):
    return ''.join(c for c in unicodedata.normalize('NFD', text)
                   if unicodedata.category(c) != 'Mn')


def keywords(description, name, key, extra):
    """Lista de sinónimos de búsqueda, sin duplicados y sin repetir el nombre."""
    parts = [clean(p) for p in re.split(r'[;\n]', description or '')]
    parts += [clean(p) for p in extra]
    seen, out = set(), []
    skip = {strip_accents(clean(name).lower()), strip_accents(clean(key).lower())}
    for part in parts:
        part = part.rstrip('.').strip()
        if not part:
            continue
        fingerprint = strip_accents(part.lower())
        if fingerprint in seen or fingerprint in skip:
            continue
        seen.add(fingerprint)
        out.append(part)
    return out


def read_dictionary():
    """Lee el diccionario maestro de búsqueda.

    Devuelve ``(por_clave, erratas_globales)``:

    * ``por_clave[clave]`` -> ``{'aliases', 'tags', 'typos'}`` del estudio.
    * ``erratas_globales`` -> ``{variante: forma_correcta}`` de la sección 4,
      que el buscador aplica a la consulta antes de puntuar.
    """
    if not os.path.exists(DICT):
        print('AVISO: falta estudios/diccionario-busqueda.md; se construye sin alias')
        return {}, {}

    text = io.open(DICT, encoding='utf-8').read()

    # El diccionario marca "sin datos" con una raya; no es un término.
    placeholders = {'—', '-', 'n/a', 'ninguno'}

    def quoted(line):
        values = (value.strip() for value in re.findall(r'`([^`]+)`', line))
        return [value for value in values if value and value.lower() not in placeholders]

    def field(block, name):
        match = re.search(r'^- \*\*' + name + r':\*\*\s*(.*)$', block, re.M)
        return quoted(match.group(1)) if match else []

    by_key = {}
    for block in re.split(r'^### \d+\.\s+', text, flags=re.M)[1:]:
        key = field(block, 'clave')
        if not key:
            continue
        typos = [t for t in field(block, 'errores_probables') if t != 'usar fuzzy global']
        by_key[key[0]] = {
            'aliases': field(block, 'aliases_directos'),
            'tags': field(block, 'terminos_relacionados'),
            'typos': typos,
        }

    # Sección 4: "- `correcta` ⇄ `variante`; `variante`; ..."
    typo_map = {}
    section = re.search(r'^## 4\..*?^## ', text, re.M | re.S)
    for line in (section.group(0) if section else '').splitlines():
        if '⇄' not in line:
            continue
        left, right = line.split('⇄', 1)
        canonical = quoted(left)
        if not canonical:
            continue
        for variant in quoted(right):
            variant = strip_accents(variant.lower())
            target = strip_accents(canonical[0].lower())
            # Una "errata" idéntica a la forma correcta no aporta nada, y una
            # de varias palabras no sirve token a token.
            if variant != target and ' ' not in variant:
                typo_map[variant] = target
    return by_key, typo_map


def build():
    dictionary, typo_map = read_dictionary()
    workbook = openpyxl.load_workbook(XLSX, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    header = [clean(h) for h in rows[0]]
    index = {name: i for i, name in enumerate(header)}

    catalog = []
    for row in rows[1:]:
        if not any(row):
            continue
        key = clean(row[index['Clave o componentes']])
        raw_category = CATEGORY_FIX.get(key, clean(row[index['Columna1']]))
        category, category_label, icon = CATEGORIES[raw_category]
        branch_code = clean(row[index['Sucursal']])
        raw_price = row[index['Precio']]
        name = clean(row[index['Nombre']])

        entry = dictionary.get(key, {})
        extra = EXTRA_KEYWORDS.get(key, []) + entry.get('aliases', [])

        item = {
            'c': key,
            'n': name,
            'cat': category,
            'catLabel': category_label,
            'icon': icon,
            'price': raw_price if isinstance(raw_price, (int, float)) else None,
            'days': int(row[index['Entrega (días)']] or 0),
            'prep': clean(row[index['Preparación para el paciente']]),
            # kw   = alias directos (puntúan alto en el buscador)
            # tags = grupos semánticos (puntúan bajo, sirven para descubrir)
            # typos = erratas concretas que el diccionario documenta
            'kw': keywords(row[index['Descripción']], name, key, extra),
            'tags': entry.get('tags', []),
            'typos': sorted({t.lower() for t in entry.get('typos', [])}),
            'branches': BRANCHES[branch_code],
            'pack': clean(row[index['Tipo']]) == 'Paquete',
            'top': key in TOP,
            'alta': key in ALTA,
            'campaign': key in CAMPAIGN,
        }
        if item['price'] is None:
            item['priceLabel'] = PRICE_LABELS.get(clean(raw_price), 'Consultar')
        catalog.append(item)

    lines = [json.dumps(item, ensure_ascii=False, separators=(',', ':')) for item in catalog]
    body = ',\n    '.join(lines)
    output = (
        '// Catálogo BIOS — generado por tools/build_estudios.py desde\n'
        '// "estudios/Top 100 estudios.xlsx" y "estudios/diccionario-busqueda.md".\n'
        '// No lo edites a mano: vuelve a correr el script cuando cambie alguno.\n'
        f'// {len(catalog)} estudios · kw = alias directos · tags = grupos\n'
        '// semánticos · typos = erratas documentadas.\n'
        'window.BIOS_ESTUDIOS = [\n    '
        + body +
        '\n];\n\n'
        '// Erratas frecuentes (sección 4 del diccionario). El buscador las\n'
        '// corrige en la consulta antes de puntuar.\n'
        'window.BIOS_TYPOS = '
        + json.dumps(typo_map, ensure_ascii=False, sort_keys=True, indent=0).replace('\n', '')
        + ';\n'
    )
    with open(OUT, 'w', encoding='utf-8') as handle:
        handle.write(output)
    print(f'{len(catalog)} estudios escritos en {os.path.relpath(OUT, ROOT)}')
    print(f'{sum(len(i["kw"]) for i in catalog)} alias directos indexados')
    print(f'{sum(len(i["tags"]) for i in catalog)} etiquetas semánticas')
    print(f'{sum(len(i["typos"]) for i in catalog)} erratas por estudio + {len(typo_map)} globales')
    missing = [i['c'] for i in catalog if not i['tags']]
    if missing:
        print(f'AVISO: sin etiquetas semánticas: {", ".join(missing)}')


if __name__ == '__main__':
    build()
